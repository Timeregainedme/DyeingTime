"""P3：DOCX / Markdown 原文提取。

约束：
- 只接受 .docx / .md；扩展名、实际字节和可解析性都要验证；
- DOCX 按正文顺序保留段落与表格，不执行宏、不访问外部引用；
- 无法提取的内容（图片等）以 warning 显式暴露，不静默丢弃正文。
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
_ALLOWED_SUFFIXES = {".docx", ".md", ".markdown"}


@dataclass(slots=True)
class ParsedScriptDocument:
    """解析结果：完整原文 + 可解释告警。"""

    file_type: str
    text: str
    warnings: list[str] = field(default_factory=list)
    table_count: int = 0
    image_count: int = 0
    char_count: int = 0


def normalize_suffix(filename: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".markdown":
        return ".md"
    return suffix


def _decode_markdown(content: bytes, filename: str) -> ParsedScriptDocument:
    if b"\x00" in content[:4096]:
        raise ValueError(f"Markdown 文件包含二进制内容，无法按文本解析：{filename}")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Markdown 文件不是有效的 UTF-8 编码") from exc
    # 统一换行，保证证据定位的字符偏移稳定；不改变可见内容。
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    warnings: list[str] = []
    if not text.strip():
        warnings.append("文档内容为空")
    return ParsedScriptDocument(
        file_type="md",
        text=text,
        warnings=warnings,
        char_count=len(text),
    )


def _heading_prefix(paragraph) -> str:
    """把 Word 标题样式转为 Markdown 井号，保留标题层级。"""
    try:
        style_name = (paragraph.style.name or "").strip().lower()
    except Exception:  # noqa: BLE001 - 非标准样式不应阻断正文提取
        return ""
    if style_name in {"title", "标题"}:
        return "# "
    if style_name.startswith("heading"):
        digits = "".join(ch for ch in style_name if ch.isdigit())
        level = int(digits) if digits else 1
        return "#" * min(max(level, 1), 6) + " "
    if style_name.startswith("标题"):
        digits = "".join(ch for ch in style_name if ch.isdigit())
        level = int(digits) if digits else 1
        return "#" * min(max(level, 1), 6) + " "
    return ""


def _parse_docx(content: bytes, filename: str) -> ParsedScriptDocument:
    if not content.startswith(b"PK\x03\x04"):
        raise ValueError("DOCX 文件头无效，可能不是真正的 Word 文档")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                raise ValueError("DOCX 缺少 word/document.xml")
    except zipfile.BadZipFile as exc:
        raise ValueError("DOCX 文件损坏或不是有效的压缩包") from exc

    try:
        from docx import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - 依赖缺失属于部署错误
        raise ValueError("服务端未安装 python-docx，无法解析 Word 文档") from exc

    try:
        document = DocxDocument(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001 - python-docx 的异常类型较多
        raise ValueError("DOCX 解析失败，请确认文件未损坏") from exc

    lines: list[str] = []
    table_count = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = (paragraph.text or "").rstrip()
            if not text:
                continue
            prefix = _heading_prefix(paragraph)
            lines.append(f"{prefix}{text}")
        elif tag == "tbl":
            table_count += 1
            table = Table(child, document)
            lines.append(f"【表格 {table_count}】")
            for row in table.rows:
                cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
            lines.append("【表格结束】")

    image_count = len(getattr(document, "inline_shapes", []) or [])
    text = "\n".join(lines).strip()
    warnings: list[str] = []
    if not text:
        warnings.append("Word 文档没有可提取的正文")
    if table_count:
        warnings.append(f"已提取 {table_count} 个表格，请核对表格内换行是否影响剧本格式")
    if image_count:
        warnings.append(f"文档包含 {image_count} 张图片，图片内容未纳入原文，请人工补充必要信息")

    return ParsedScriptDocument(
        file_type="docx",
        text=text,
        warnings=warnings,
        table_count=table_count,
        image_count=image_count,
        char_count=len(text),
    )


def parse_script_document(*, filename: str, content: bytes) -> ParsedScriptDocument:
    """按扩展名分派解析；扩展名与实际内容不匹配时直接报错。"""
    if not filename:
        raise ValueError("上传文件缺少文件名")
    suffix = normalize_suffix(filename)
    if suffix not in _ALLOWED_SUFFIXES:
        raise ValueError("只支持 .docx 与 .md 剧本文件")
    if not content:
        raise ValueError("上传文件为空")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError(f"文件超过大小限制（{MAX_DOCUMENT_BYTES // 1024 // 1024}MB）")

    if suffix == ".docx":
        return _parse_docx(content, filename)
    return _decode_markdown(content, filename)


__all__ = [
    "MAX_DOCUMENT_BYTES",
    "ParsedScriptDocument",
    "normalize_suffix",
    "parse_script_document",
]
