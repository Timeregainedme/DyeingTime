"""P3 / C5：剧本文档解析与版本管理测试。"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from starlette.datastructures import UploadFile

from app.core.db import Base
from app.core.storage import StoredFileInfo
from app.models.studio import Chapter, Project, ProjectStyle, ProjectVisualStyle
from app.services.studio import script_documents as svc
from app.services.studio.script_document_parser import parse_script_document


def _docx_bytes(*, with_table: bool = False) -> bytes:
    document = DocxDocument()
    document.add_heading("第一集：雨夜", level=1)
    document.add_paragraph("场景：出租屋，夜，内景。")
    document.add_paragraph("林晚：你终于回来了。")
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "角色"
        table.cell(0, 1).text = "服装"
        table.cell(1, 0).text = "林晚"
        table.cell(1, 1).text = "米色风衣"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _build_session() -> tuple[AsyncSession, object]:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", future=True)
    session_local = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    return session_local(), engine


async def _seed(db: AsyncSession) -> None:
    db.add(
        Project(
            workspace_id="ws-test",
            id="p1",
            name="测试项目",
            description="",
            style=ProjectStyle.real_people_city,
            visual_style=ProjectVisualStyle.live_action,
        )
    )
    db.add(
        Chapter(
            id="c1",
            project_id="p1",
            index=1,
            title="第1集：雨夜",
            raw_text="旧正文",
        )
    )
    await db.commit()


def _upload(name: str, content: bytes) -> UploadFile:
    return UploadFile(file=io.BytesIO(content), filename=name)


@pytest.fixture(autouse=True)
def _mock_storage(monkeypatch):
    async def _fake_upload(*, key, data, content_type=None, extra_args=None):
        return StoredFileInfo(key=key, url=f"https://r2.example/{key}", size=len(data))

    monkeypatch.setattr(svc.storage, "upload_file", _fake_upload)
    return _fake_upload


def test_markdown_parser_preserves_headings_and_unicode() -> None:
    content = "# 第一集：雨夜\n\n## 场景 1\n林晚：你好。\n\n| 角色 | 服装 |\n| --- | --- |\n| 林晚 | 风衣 |\n".encode()
    parsed = parse_script_document(filename="第一集：雨夜.md", content=content)
    assert parsed.file_type == "md"
    assert "## 场景 1" in parsed.text
    assert "林晚：你好。" in parsed.text
    assert parsed.char_count == len(parsed.text)


def test_docx_parser_preserves_paragraph_order_and_tables() -> None:
    parsed = parse_script_document(filename="雨夜.docx", content=_docx_bytes(with_table=True))
    assert parsed.file_type == "docx"
    assert parsed.table_count == 1
    assert parsed.text.startswith("# 第一集：雨夜")
    assert "场景：出租屋" in parsed.text
    assert "林晚 | 米色风衣" in parsed.text
    assert parsed.image_count == 0


def test_parser_rejects_invalid_formats_and_marks_empty_docs() -> None:
    with pytest.raises(ValueError):
        parse_script_document(filename="剧本.pdf", content=b"pdf")
    empty = parse_script_document(filename="空.md", content=b"   \n")
    assert empty.text.strip() == ""
    assert any("为空" in warning for warning in empty.warnings)
    with pytest.raises(ValueError):
        parse_script_document(filename="假.docx", content=b"not-a-zip")
    with pytest.raises(ValueError):
        parse_script_document(filename="坏.md", content=b"\x00\x01")


@pytest.mark.asyncio
async def test_upload_creates_version_and_duplicate_is_idempotent() -> None:
    db, engine = await _build_session()
    async with db:
        await _seed(db)
        project = await db.get(Project, "p1")
        chapter = await db.get(Chapter, "c1")
        assert project is not None and chapter is not None

        docx_content = _docx_bytes()
        first = await svc.upload_script_document(
            db,
            project=project,
            chapter=chapter,
            doc_type=svc.DOC_TYPE_EPISODE,
            file=_upload("第1集.docx", docx_content),
            created_by=1,
        )
        assert first.created is True
        assert first.version.version_no == 1
        assert first.version.status == "candidate"
        assert first.document.active_version_id is None
        # 上传阶段不覆盖正式章节正文
        assert chapter.raw_text == "旧正文"

        duplicate = await svc.upload_script_document(
            db,
            project=project,
            chapter=chapter,
            doc_type=svc.DOC_TYPE_EPISODE,
            file=_upload("第1集-副本.docx", docx_content),
            created_by=1,
        )
        assert duplicate.created is False
        assert duplicate.version.id == first.version.id
        assert len(duplicate.document.versions) == 1
    await engine.dispose()


@pytest.mark.asyncio
async def test_activate_version_syncs_chapter_and_keeps_history() -> None:
    db, engine = await _build_session()
    async with db:
        await _seed(db)
        project = await db.get(Project, "p1")
        chapter = await db.get(Chapter, "c1")
        assert project is not None and chapter is not None

        first = await svc.upload_script_document(
            db,
            project=project,
            chapter=chapter,
            doc_type=svc.DOC_TYPE_EPISODE,
            file=_upload("第1集.md", "# 第一版\n林晚：旧台词。".encode()),
        )
        await svc.activate_version(
            db, document_id=first.document.id, version_id=first.version.id
        )
        await db.refresh(chapter)
        assert chapter.raw_text == "# 第一版\n林晚：旧台词。"

        second = await svc.upload_script_document(
            db,
            project=project,
            chapter=chapter,
            doc_type=svc.DOC_TYPE_EPISODE,
            file=_upload("第1集-改.md", "# 第二版\n林晚：新台词。".encode()),
        )
        assert second.version.status == "candidate"
        assert second.document.active_version_id == first.version.id

        activated = await svc.activate_version(
            db, document_id=second.document.id, version_id=second.version.id
        )
        await db.refresh(chapter)
        assert chapter.raw_text == "# 第二版\n林晚：新台词。"
        statuses = {item.id: item.status for item in activated.versions}
        assert statuses[first.version.id] == "superseded"
        assert statuses[second.version.id] == "active"
    await engine.dispose()


@pytest.mark.asyncio
async def test_worldbuilding_document_has_no_chapter_and_versions_increment() -> None:
    db, engine = await _build_session()
    async with db:
        await _seed(db)
        project = await db.get(Project, "p1")
        assert project is not None

        first = await svc.upload_script_document(
            db,
            project=project,
            doc_type=svc.DOC_TYPE_WORLD,
            file=_upload("世界观.md", "# 世界观\n现代都市。".encode()),
        )
        second = await svc.upload_script_document(
            db,
            project=project,
            doc_type=svc.DOC_TYPE_WORLD,
            file=_upload("世界观-v2.md", "# 世界观 v2\n现代都市，雨夜。".encode()),
        )
        assert first.document.chapter_id is None
        assert second.version.version_no == 2
        documents = await svc.list_project_documents(db, project_id="p1")
        assert len(documents.items) == 1
        assert len(documents.items[0].versions) == 2
    await engine.dispose()


@pytest.mark.asyncio
async def test_upload_rejects_wrong_scope() -> None:
    db, engine = await _build_session()
    async with db:
        await _seed(db)
        project = await db.get(Project, "p1")
        chapter = await db.get(Chapter, "c1")
        assert project is not None and chapter is not None

        with pytest.raises(ValueError):
            await svc.upload_script_document(
                db,
                project=project,
                doc_type=svc.DOC_TYPE_EPISODE,
                file=_upload("第1集.md", "# x".encode()),
            )
        with pytest.raises(ValueError):
            await svc.upload_script_document(
                db,
                project=project,
                chapter=chapter,
                doc_type=svc.DOC_TYPE_WORLD,
                file=_upload("世界观.md", "# x".encode()),
            )
    await engine.dispose()


def test_docx_zip_without_document_xml_is_rejected() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/other.xml", "x")
    with pytest.raises(ValueError):
        parse_script_document(filename="坏.docx", content=buffer.getvalue())
